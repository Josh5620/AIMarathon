"""
Resume text extraction (PDF + DOCX).

Scanned-PDF OCR fallback uses Tesseract via pytesseract. Install:
  Ubuntu/VPS:  sudo apt-get install tesseract-ocr
  Windows dev: https://github.com/UB-Mannheim/tesseract/wiki
If the binary is missing, OCR silently degrades — direct-extraction paths
are unaffected and the guard will mark text-less PDFs as needs_review.
"""
import io
from pathlib import Path

import fitz  # PyMuPDF
import docx

try:
    import pytesseract
    from PIL import Image
    _OCR_AVAILABLE = True
except ImportError:
    _OCR_AVAILABLE = False

ALLOWED_EXTENSIONS = (".pdf", ".docx")

# Cascade threshold: PDF direct-extraction text below this char count triggers OCR.
# Per 2ndHandover.md L16: "fall back to OCR only if you get < 100 chars back".
OCR_MIN_CHARS = 100

# Rasterisation resolution for Tesseract. 200 dpi balances accuracy and speed.
OCR_DPI = 200


class EncryptedFileError(Exception):
    """Raised when a PDF is password-protected and cannot be read."""


def extract_text(file_bytes: bytes, filename: str) -> str:
    """
    Extract plain text from a PDF or DOCX file.

    Raises:
        ValueError: if the file extension is not .pdf or .docx.
        EncryptedFileError: if the PDF is password-protected.
    """
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"unsupported file type: '{ext}' — only .pdf and .docx are accepted")

    if ext == ".docx":
        return _extract_docx(file_bytes)
    return _extract_pdf(file_bytes)


def _extract_docx(file_bytes: bytes) -> str:
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception:
        return ""  # corrupt DOCX — ingest_resume will treat empty as needs_review

    parts: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    parts.append(text)

    return "\n".join(parts)


def _extract_pdf(file_bytes: bytes) -> str:
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception:
        return ""  # corrupt PDF bytes

    if doc.needs_pass or doc.is_encrypted:
        raise EncryptedFileError("password-protected PDF — cannot extract text")

    text = "\n".join(page.get_text("text") for page in doc)

    # Cascade: if direct extraction yielded very little text the PDF is likely
    # a scanned image. Try Tesseract OCR as a fallback (per 2ndHandover.md L16).
    if len(text.strip()) < OCR_MIN_CHARS:
        ocr_text = _ocr_pdf(doc)
        if ocr_text.strip():
            return ocr_text

    return text


def _ocr_pdf(doc: fitz.Document) -> str:
    """
    Rasterize each page and run Tesseract OCR.

    Returns '' if pytesseract is not importable or the Tesseract binary is not
    installed — the guard will then mark the result needs_review.
    """
    if not _OCR_AVAILABLE:
        return ""

    pages: list[str] = []
    for page in doc:
        try:
            pix = page.get_pixmap(dpi=OCR_DPI)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            pages.append(pytesseract.image_to_string(img))
        except pytesseract.TesseractNotFoundError:
            # Tesseract binary not installed → bail out of the entire cascade
            return ""
        except Exception:
            # Per-page failure (corrupt page, encoding issue) — skip and continue
            continue

    return "\n".join(pages)
