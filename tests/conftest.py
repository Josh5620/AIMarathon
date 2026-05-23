"""Shared pytest fixtures for the ingestion test suite."""
from io import BytesIO
from pathlib import Path

import fitz  # PyMuPDF
import docx
import pytest

import app.ingestion.audit as audit_module


@pytest.fixture()
def tmp_audit_log(tmp_path, monkeypatch):
    """Redirect audit log writes to a temp file so tests don't pollute audit.log."""
    log_path = tmp_path / "audit.log"
    monkeypatch.setattr(audit_module, "AUDIT_LOG_PATH", log_path)
    return log_path


def make_pdf(text: str, password: str | None = None) -> bytes:
    """Build a one-page text PDF in memory, optionally encrypted."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontsize=11)
    if password:
        doc.save(
            BytesIO(),
            encryption=fitz.PDF_ENCRYPT_AES_256,
            user_pw=password,
            owner_pw=password,
        )
        buf = BytesIO()
        doc.save(
            buf,
            encryption=fitz.PDF_ENCRYPT_AES_256,
            user_pw=password,
            owner_pw=password,
        )
        return buf.getvalue()
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_docx(text: str, table_text: str | None = None) -> bytes:
    """Build a DOCX in memory, optionally with a one-cell table."""
    doc = docx.Document()
    doc.add_paragraph(text)
    if table_text:
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = table_text
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture()
def sample_pdf():
    return make_pdf("Alice Smith\nalice@example.com\nSenior Python engineer, 8 years.")


@pytest.fixture()
def sample_docx():
    return make_docx(
        "Bob Jones\nbob@example.com\nJava developer, 5 years.",
        table_text="Skills: Java, Spring Boot",
    )
